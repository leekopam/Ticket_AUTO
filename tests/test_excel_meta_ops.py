"""ExcelService 내부 시트/원자적 저장/작업 이력 테스트."""
from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from services.excel_service import (
    META_SHEET,
    OPERATIONS_SHEET,
    ExcelService,
)


def _make_orders_xlsx(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "주문목록"
    ws.append(["주문번호", "주문자명", "주문자연락처", "좌석번호", "주문상태", "[상품1]티켓"])
    ws.append(["AAAA1111_BBBB2222", "홍길동", "010-1234-5678", "A-1", "결제완료", 1])
    ws.append(["CCCC3333_DDDD4444", "김철수", "010-9999-8888", "A-2", "주문취소", 2])
    wb.save(path)


@pytest.fixture
def excel(tmp_path: Path) -> ExcelService:
    data = tmp_path / "data.xlsx"
    _make_orders_xlsx(data)
    return ExcelService(str(data))


def test_atomic_save_creates_backup_and_valid_file(excel: ExcelService, tmp_path: Path):
    target = tmp_path / "data.xlsx"
    assert excel.mark_order_received("AAAA1111_BBBB2222", "2026-01-01 10:00:00")

    backup = tmp_path / "data.xlsx.bak"
    assert backup.exists(), "백업 파일이 없음"
    assert not (tmp_path / "data.xlsx.tmp.xlsx").exists(), "임시 파일이 남아있음"

    check = load_workbook(target)
    try:
        ws = check["주문목록"]
        assert ws.cell(row=2, column=6).value == 1
    finally:
        check.close()

    order = excel.find_order("AAAA1111_BBBB2222")
    assert order is not None and order.received_at == "2026-01-01 10:00:00"


def test_dataset_id_stable_across_writes(excel: ExcelService):
    first = excel.ensure_dataset_id()
    assert first
    assert excel.ensure_dataset_id() == first
    excel.mark_order_received("AAAA1111_BBBB2222", "2026-01-01 10:00:00")
    # 쓰기는 세대를 바꾸지 않는다
    assert excel.ensure_dataset_id() == first


def test_operations_roundtrip(excel: ExcelService):
    record = {
        "request_id": "req-1",
        "order_id": "AAAA1111_BBBB2222",
        "action": "receipt",
        "device_id": "phone-1",
        "state": "accepted",
        "created_at": "2026-01-01 10:00:00",
        "updated_at": "2026-01-01 10:00:00",
    }
    assert excel.append_operation(record)

    fetched = excel.get_operation("req-1")
    assert fetched is not None
    assert fetched["order_id"] == "AAAA1111_BBBB2222"
    assert fetched["state"] == "accepted"

    assert excel.update_operation("req-1", {"state": "succeeded"})
    fetched = excel.get_operation("req-1")
    assert fetched["state"] == "succeeded"
    assert fetched["updated_at"]

    assert excel.get_operation("missing") is None
    assert len(excel.list_operations()) == 1


def test_internal_sheets_do_not_break_order_lookup(excel: ExcelService):
    excel.ensure_dataset_id()
    excel.append_operation({"request_id": "r1", "order_id": "o1", "action": "receipt"})

    wb = load_workbook(excel._file_path)
    try:
        assert META_SHEET in wb.sheetnames
        assert OPERATIONS_SHEET in wb.sheetnames
        assert wb[META_SHEET].sheet_state == "hidden"
        assert wb[OPERATIONS_SHEET].sheet_state == "hidden"
    finally:
        wb.close()

    # 내부 시트가 있어도 데이터 시트를 올바르게 선택해야 한다
    order = excel.find_order("AAAA1111_BBBB2222")
    assert order is not None and order.name == "홍길동"
    assert len(excel.search_orders_all("")) == 2
    assert len(excel.search_orders("")) == 2


def test_search_orders_all_has_no_200_cap(excel: ExcelService, tmp_path: Path):
    path = tmp_path / "big.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "주문목록"
    ws.append(["주문번호", "주문자명"])
    for i in range(250):
        ws.append([f"ORD{i:04d}_X{i:012d}", f"고객{i}"])
    wb.save(path)

    big = ExcelService(str(path))
    assert len(big.search_orders("")) == 200
    assert len(big.search_orders_all("")) == 250


def test_data_file_signature_changes_on_write(excel: ExcelService):
    before = excel.data_file_signature()
    excel.mark_order_received("AAAA1111_BBBB2222", "2026-01-01 10:00:00")
    after = excel.data_file_signature()
    assert before and after and before != after
