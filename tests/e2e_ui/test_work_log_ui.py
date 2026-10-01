"""U05 티켓 업무 탭 UI 시나리오 (Flet web + headless Playwright).

워크북에 수령확인을 직접 기록해 mtime watcher 자동 갱신까지 검증한다:
탭 진입 → 빈 상태 → 수령 기록 → 목록 표시 → 행 선택 → 상세 표시.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from e2e.support import TEST_ORDER_NUMBER
from e2e_ui.support import wait_for_button

_TIMEOUT_MS = 20000


def _mark_order_received(
    data_path: Path,
    order_number: str,
    timestamp: str,
    processing_time: str = "",
) -> None:
    """테스트 워크북의 수령확인 셀에 타임스탬프를 기록한다. processing_time이 주어지면 처리시간 셀도 함께 쓴다."""
    workbook = load_workbook(data_path)
    try:
        ws = workbook.active
        headers = {str(cell.value or "").strip(): idx for idx, cell in enumerate(ws[1], 1)}
        order_col = headers.get("주문번호")
        received_col = headers.get("수령확인")
        processing_col = headers.get("처리시간")
        assert order_col and received_col, "테스트 워크북에 필요한 헤더가 없습니다."
        for row in range(2, ws.max_row + 1):
            if str(ws.cell(row=row, column=order_col).value or "").strip() == order_number:
                ws.cell(row=row, column=received_col, value=timestamp)
                if processing_time and processing_col:
                    ws.cell(row=row, column=processing_col, value=processing_time)
                workbook.save(data_path)
                return
        raise AssertionError(f"테스트 워크북에 주문 {order_number}가 없습니다.")
    finally:
        workbook.close()


def _open_work_log_tab(page) -> None:
    wait_for_button(page, "처리 현황 조회", timeout_ms=_TIMEOUT_MS).click()


def test_work_log_tab_empty_then_row_and_detail(page, flet_server):
    _open_work_log_tab(page)
    page.get_by_text("아직 처리된 주문이 없습니다.", exact=True).first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )

    data_path = Path(flet_server["runtime_dir"]) / "Resources" / "data" / "data.xlsx"
    # 수령확인(재확인 시각)과 다른 처리시간을 기록해 처리시간 컬럼이 표시 기준인지 검증한다.
    _mark_order_received(
        data_path,
        TEST_ORDER_NUMBER,
        "2026-09-27 18:45:00",
        processing_time="2026-09-27 12:00:00",
    )

    # mtime watcher(0.5s)가 파일 변경을 감지해 목록을 자동 갱신한다.
    # 행은 클릭 가능한 컨테이너라 텍스트가 부모 semantics 노드로 병합되므로 부분 매칭을 사용한다.
    page.get_by_text("테스트 사용자").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    page.get_by_text("처리 완료 손님").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    # 상태 컬럼은 제거됨 — 수령완료 배지가 나오면 안 된다
    assert page.get_by_text("수령완료").count() == 0
    # 티켓/상품 이름은 클릭 없이 목록 행에 바로 표시된다 (행 병합 라벨 기준 부분 매칭)
    page.get_by_text("테스트 상품").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    # PC 본체 처리 건은 행 우측에 처리 단말 칩 "PC"가 표시된다
    # (행이 클릭 가능 컨테이너라 자식 텍스트가 부모 노드로 병합되므로 부분 매칭)
    page.get_by_text("PC").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    # 처리시간 열은 data.xlsx의 처리시간 컬럼 값(MM-DD HH:MM:SS)을 표시한다
    page.get_by_text("09-27 12:00:00").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    assert page.get_by_text("09-27 18:45:00").count() == 0

    # 행 선택 → 우측 상세에 주문 정보와 전달 상품이 표시된다.
    row = page.get_by_text("테스트 사용자").first
    row.click()
    page.get_by_text("상세정보").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    page.get_by_text("테스트 상품").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    page.get_by_text("010-0000-0000").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )


def test_work_log_reset_button_clears_processed_state(page, flet_server):
    """처리 데이터 초기화 버튼이 확인 다이얼로그를 거쳐 수령 처리 표시를 비운다."""
    data_path = Path(flet_server["runtime_dir"]) / "Resources" / "data" / "data.xlsx"
    _mark_order_received(data_path, TEST_ORDER_NUMBER, "2026-09-27 18:45:00")

    _open_work_log_tab(page)
    page.get_by_text("테스트 사용자").first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )

    # 초기화 버튼 → 확인 다이얼로그 → 초기화 확정
    wait_for_button(page, "처리 데이터 초기화", timeout_ms=_TIMEOUT_MS).click()
    page.get_by_text("수령 처리 표시", exact=False).first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    wait_for_button(page, "초기화", timeout_ms=_TIMEOUT_MS).click()

    # 목록이 빈 상태로 돌아가고 주문 데이터 자체는 남는다
    page.get_by_text("아직 처리된 주문이 없습니다.", exact=True).first.wait_for(
        state="visible", timeout=_TIMEOUT_MS
    )
    workbook = load_workbook(data_path, read_only=True, data_only=True)
    try:
        headers = {str(c.value or "").strip(): i for i, c in enumerate(workbook.active[1], 1)}
        order_col, received_col = headers["주문번호"], headers["수령확인"]
        remaining = [
            row for row in workbook.active.iter_rows(min_row=2, values_only=True)
            if str(row[received_col - 1] or "").strip()
        ]
        assert remaining == []
    finally:
        workbook.close()
