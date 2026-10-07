"""계획서 §2의 S01~S08 시나리오를 오프라인 대역으로 검증한다."""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import main as app_main
from services.api_service import ApiService
from services.browser_service import (
    BrowserResolveResult,
    PageOrderDiscoveryResult,
    ReceiptClickResult,
)
from services.excel_service import ExcelService

from .support import (
    TEST_ORDER_NUMBER,
    TEST_QR_URL,
    FakeBrowserService,
    FakePrinterBackend,
    FakeScannerView,
    build_offline_app,
    create_test_workbook,
    decode_generated_qr,
)


def _build_app(
    data_path: Path,
    *,
    offline_scan_mode: bool = True,
    browser: FakeBrowserService | None = None,
    scanner: FakeScannerView | None = None,
):
    """공용 오프라인 앱 조립을 재사용한다."""
    return build_offline_app(
        data_path,
        offline_scan_mode=offline_scan_mode,
        browser=browser,
        scanner=scanner,
    )


def _resolved_detail_result(order_number: str = TEST_ORDER_NUMBER) -> BrowserResolveResult:
    return BrowserResolveResult(
        ok=True,
        status_code=302,
        location=f"/w/myform/sellForm-history-detail/{order_number}",
    )


def _read_order_cell(data_path: Path, order_number: str, header: str) -> str:
    """워크북에서 주문번호 행의 지정 헤더 셀 값을 읽는다."""
    workbook = load_workbook(data_path, read_only=True, data_only=True)
    try:
        ws = workbook.active
        headers = {str(cell.value or "").strip(): idx for idx, cell in enumerate(ws[1], 1)}
        order_col = headers.get("주문번호")
        target_col = headers.get(header)
        assert order_col and target_col, f"테스트 워크북에 필요한 헤더가 없습니다: {header}"
        for row in ws.iter_rows(min_row=2, values_only=True):
            if str(row[order_col - 1] or "").strip() == order_number:
                return str(row[target_col - 1] or "").strip()
        raise AssertionError(f"테스트 워크북에 주문 {order_number}가 없습니다.")
    finally:
        workbook.close()


class S01OfflineModeEntryTest(unittest.TestCase):
    """S01: 오프라인 스캔 모드 진입/스캔/모드 OFF 복귀."""

    def test_offline_entry_ready_then_scan_then_off_returns_to_auth_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            app, browser, scanner, sound = _build_app(data_path)
            printer = FakePrinterBackend()

            # 오프라인 모드 진입: READY + auth_ready/scanning 활성화 (QR §1-1)
            app._enter_ready("[오프라인] 로그인 없이 스캔 테스트 모드")
            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertTrue(scanner.auth_ready)
            self.assertTrue(scanner.scanning_enabled)
            self.assertEqual(scanner.status_message, "[오프라인] 로그인 없이 스캔 테스트 모드")

            with (
                patch("httpx.get") as http_get,
                patch(
                    "services.receipt_print_pipeline.WindowsPrinterService",
                    return_value=printer,
                ),
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            http_get.assert_not_called()
            self.assertEqual(sound.success_count, 1)
            self.assertEqual(len(printer.jobs), 1)

            # 오프라인 스캔도 수령확인과 처리시간을 모두 data.xlsx에 기록한다
            received_at = _read_order_cell(data_path, TEST_ORDER_NUMBER, "수령확인")
            processing_time = _read_order_cell(data_path, TEST_ORDER_NUMBER, "처리시간")
            self.assertTrue(received_at)
            self.assertEqual(processing_time, received_at)

            # 오프라인 모드 OFF: 브라우저 경로로 전환되어 AUTH_REQUIRED → 복구 대기 (QR §1-4)
            app._ticket_debug_tools_service.settings.offline_scan_mode = False
            browser.resolve_results.append(
                BrowserResolveResult(ok=False, error_code="AUTH_REQUIRED")
            )
            with patch.object(app, "_wait_for_login", return_value=False):
                app._process_qr(TEST_QR_URL, allow_auth_retry=True)

            self.assertIn(("resolve_qr_redirect", TEST_QR_URL), browser.calls)
            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertEqual(scanner.status_message, "로그인 대기 시간 초과 - 준비 상태로 돌아갑니다.")


class S02OrderNumberRecoveryTest(unittest.TestCase):
    """S02: QR 파싱 실패 시 주문번호 복구 3경로."""

    def _process_missing_order(self, data_path: Path, discovery: PageOrderDiscoveryResult):
        browser = FakeBrowserService()
        browser.discovery = discovery
        app, browser, scanner, sound = _build_app(data_path, browser=browser)
        printer = FakePrinterBackend()
        # 주문 상세 경로이나 주문번호 추출 불가 → ORDER_NUMBER_MISSING
        browser_result = BrowserResolveResult(
            ok=True,
            status_code=302,
            location="/w/myform/sellForm-history-detail/",
        )
        with patch(
            "services.receipt_print_pipeline.WindowsPrinterService",
            return_value=printer,
        ):
            app._process_resolved_qr(TEST_QR_URL, browser_result, allow_auth_retry=False)
        return app, scanner, sound, printer

    def test_page_order_number_discovery_recovers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            discovery = PageOrderDiscoveryResult(
                order_number=TEST_ORDER_NUMBER,
                url=f"{ApiService.WITCHFORM_BASE}/w/myform/sellForm-history-detail/{TEST_ORDER_NUMBER}",
            )
            app, scanner, sound, printer = self._process_missing_order(data_path, discovery)

            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertEqual(scanner.status_message, "수령 완료 및 영수증 출력 완료")
            self.assertTrue(ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER).is_received)

    def test_customer_context_excel_match_recovers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            # 페이지에서 주문번호는 못 찾고 고객 정보만 발견 → 엑셀 매칭
            discovery = PageOrderDiscoveryResult(
                buyer_name="테스트 사용자",
                buyer_phone="010-0000-0000",
            )
            app, scanner, sound, printer = self._process_missing_order(data_path, discovery)

            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertTrue(ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER).is_received)

    def test_all_recovery_paths_fail_enters_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            discovery = PageOrderDiscoveryResult()  # 아무 정보도 없음
            app, scanner, sound, printer = self._process_missing_order(data_path, discovery)

            self.assertEqual(app._state, app_main.AppState.ERROR)
            self.assertFalse(ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER).is_received)
            self.assertEqual(len(printer.jobs), 0)


class S03ReceiptClickRetryTest(unittest.TestCase):
    """S03: 수령완료 클릭 실패 시 1.5초 후 자동 1회 재시도."""

    def test_primary_click_fail_retries_once_then_succeeds(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            browser = FakeBrowserService()
            browser.resolve_results.append(_resolved_detail_result())
            browser.click_results.extend(
                [
                    ReceiptClickResult(success=False, error_code="PRIMARY_CLICK_FAIL"),
                    ReceiptClickResult(success=True),
                ]
            )
            app, browser, scanner, sound = _build_app(
                data_path, offline_scan_mode=False, browser=browser
            )
            printer = FakePrinterBackend()

            with patch(
                "services.receipt_print_pipeline.WindowsPrinterService",
                return_value=printer,
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            # 재시도 포함 2회 클릭 → 수령 완료 + 카운트
            click_calls = [c for c in browser.calls if c[0] == "click_receipt_button"]
            self.assertEqual(len(click_calls), 2)
            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertEqual(scanner.status_message, "수령 완료 및 영수증 출력 완료")
            saved_order = ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER)
            self.assertTrue(saved_order.is_received)
            self.assertEqual(sound.success_count, 1)

    def test_retry_failure_keeps_received_mark_uncommitted(self) -> None:
        """재시도도 실패하면 수령표시·카운트 모두 미실행 (QR §3 불변규칙)."""
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            browser = FakeBrowserService()
            browser.resolve_results.append(_resolved_detail_result())
            browser.click_results.extend(
                [
                    ReceiptClickResult(success=False, error_code="PRIMARY_CLICK_FAIL"),
                    ReceiptClickResult(success=False, error_code="CONFIRM_CLICK_FAIL"),
                ]
            )
            app, browser, scanner, sound = _build_app(
                data_path, offline_scan_mode=False, browser=browser
            )
            printer = FakePrinterBackend()

            with patch(
                "services.receipt_print_pipeline.WindowsPrinterService",
                return_value=printer,
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            self.assertEqual(app._state, app_main.AppState.ERROR)
            saved_order = ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER)
            self.assertFalse(saved_order.is_received)
            self.assertEqual(sound.success_count, 0)
            self.assertEqual(len(printer.jobs), 0)


class S04PrintFailureRollbackTest(unittest.TestCase):
    """S04: 영수증 출력 실패 시 수령확인·주문상태 롤백 후 재스캔 재처리."""

    def test_print_failure_rolls_back_and_rescan_recovers(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            app, browser, scanner, sound = _build_app(data_path)
            printer = FakePrinterBackend(failure=RuntimeError("가상 프린터 실패"))

            # 출력 실패 → 수령확인·주문상태 원복 (QR §4-1)
            with patch(
                "services.receipt_print_pipeline.WindowsPrinterService",
                return_value=printer,
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            self.assertEqual(app._state, app_main.AppState.ERROR)
            self.assertIn("원복", scanner.status_message)
            failed_order = ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER)
            self.assertFalse(failed_order.is_received)
            self.assertEqual(failed_order.order_status, "거래중")
            self.assertEqual(sound.success_count, 0)
            self.assertEqual(len(printer.jobs), 1)

            # 롤백 후 재스캔 → 정상 재처리 (QR §4-2)
            printer.failure = None
            with patch(
                "services.receipt_print_pipeline.WindowsPrinterService",
                return_value=printer,
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertEqual(scanner.status_message, "수령 완료 및 영수증 출력 완료")
            recovered_order = ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER)
            self.assertTrue(recovered_order.is_received)
            self.assertEqual(recovered_order.order_status, "거래종료")
            self.assertEqual(sound.success_count, 1)
            self.assertEqual(len(printer.jobs), 2)


class S05SessionTimeoutRecoveryTest(unittest.TestCase):
    """S05: 로그인 대기 타임아웃 시 RECOVERING → READY 복귀 후 재스캔 가능."""

    def test_timeout_returns_to_ready_and_rescan_works(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            app, browser, scanner, sound = _build_app(data_path)
            app._state = app_main.AppState.RECOVERING
            scanner.auth_ready = False

            # 로그인 대기 180초 초과 재현 → READY 복귀 (QR §5-1)
            with patch.object(app, "_wait_for_login", return_value=False):
                app._recover_auth_and_retry(TEST_QR_URL)

            self.assertEqual(app._state, app_main.AppState.READY)
            self.assertTrue(scanner.scanning_enabled)
            self.assertEqual(scanner.status_message, "로그인 대기 시간 초과 - 준비 상태로 돌아갑니다.")

            # READY 복귀 후 재스캔 정상 처리 (QR §5-2)
            printer = FakePrinterBackend()
            with (
                patch("httpx.get") as http_get,
                patch(
                    "services.receipt_print_pipeline.WindowsPrinterService",
                    return_value=printer,
                ),
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            http_get.assert_not_called()
            self.assertEqual(sound.success_count, 1)
            self.assertEqual(scanner.status_message, "수령 완료 및 영수증 출력 완료")


class S06OfflineQrResolutionTest(unittest.TestCase):
    """S06: httpx 오프라인 QR 해석 — 로그인 리다이렉트 추적과 fast path."""

    def test_login_redirect_chain_extracts_order_from_final_url(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            app, browser, scanner, sound = _build_app(data_path)
            printer = FakePrinterBackend()

            # 실제 witchform QR → 최종 도착지가 로그인 페이지 (QR §6-1, §6-2)
            final_login_url = (
                f"https://witchform.com/w/login?redirect="
                f"%2Fw%2Fmyform%2FsellForm-history-detail%2F{TEST_ORDER_NUMBER}"
            )
            fake_response = unittest.mock.Mock()
            fake_response.status_code = 200
            fake_response.url = final_login_url

            with (
                patch("httpx.get", return_value=fake_response) as http_get,
                patch(
                    "services.receipt_print_pipeline.WindowsPrinterService",
                    return_value=printer,
                ),
            ):
                app._process_qr(
                    "https://witchform.com/qrcode_link.php?code=REAL_QR",
                    allow_auth_retry=False,
                )

            http_get.assert_called_once()
            self.assertTrue(ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER).is_received)
            self.assertEqual(app._state, app_main.AppState.READY)

    def test_test_order_param_skips_http_fast_path(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            app, *_ = _build_app(data_path)
            printer = FakePrinterBackend()

            with (
                patch("httpx.get") as http_get,
                patch(
                    "services.receipt_print_pipeline.WindowsPrinterService",
                    return_value=printer,
                ),
            ):
                app._process_qr(TEST_QR_URL, allow_auth_retry=False)

            http_get.assert_not_called()
            self.assertTrue(ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER).is_received)


class S07QrGenerateScanRoundTripTest(unittest.TestCase):
    """S07: 개발자도구 QR 생성 → 디코딩 → 오프라인 스캔 왕복."""

    def test_generated_qr_decodes_and_processes_order(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            app, browser, scanner, sound = _build_app(data_path)
            printer = FakePrinterBackend()

            payload = TEST_QR_URL
            decoded = decode_generated_qr(payload)
            self.assertEqual(decoded, payload)

            with (
                patch("httpx.get") as http_get,
                patch(
                    "services.receipt_print_pipeline.WindowsPrinterService",
                    return_value=printer,
                ),
            ):
                app._process_qr(decoded or "", allow_auth_retry=False)

            http_get.assert_not_called()
            self.assertEqual(scanner.status_message, "수령 완료 및 영수증 출력 완료")
            self.assertTrue(ExcelService(str(data_path)).find_order(TEST_ORDER_NUMBER).is_received)


class S08DuplicateQrTest(unittest.TestCase):
    """S08: 동일 QR 연속 스캔 시 쿨다운 내 중복은 무시한다."""

    def test_duplicate_scan_within_cooldown_is_ignored(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            data_path = Path(temp_dir) / "test_orders.xlsx"
            create_test_workbook(data_path)
            scanner = FakeScannerView()
            app, browser, scanner, sound = _build_app(data_path, scanner=scanner)
            printer = FakePrinterBackend()
            app._recent_qr = {}
            app._qr_repeat_cooldown_sec = 2.0

            scanner.push_qr(TEST_QR_URL)
            scanner.push_qr(TEST_QR_URL)  # 2초 쿨다운 내 동일 QR

            with (
                patch("httpx.get") as http_get,
                patch(
                    "services.receipt_print_pipeline.WindowsPrinterService",
                    return_value=printer,
                ),
            ):
                scanner.start()
                app._main_loop()

            http_get.assert_not_called()
            # 첫 스캔만 처리되고 두 번째는 무시됨
            self.assertEqual(sound.success_count, 1)
            self.assertEqual(len(printer.jobs), 1)
            self.assertIn("중복 QR 코드 무시됨", scanner.status_history)


if __name__ == "__main__":
    unittest.main()
